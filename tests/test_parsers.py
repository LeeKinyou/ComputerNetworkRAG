from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.parsers import get_parser
from app.parsers.base import (
    BaseParser,
    NotEnabledError,
    PARSER_REGISTRY,
    ParsedDocument,
    UnsupportedFormatError,
)

MD_SAMPLE = """# 第5章 传输层

传输层提供**端到端**的通信服务。

## 5.1 UDP 协议

UDP 是无连接的传输层协议，首部仅 8 字节。
"""


def _make_sample_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("第4章 网络层", level=1)
    doc.add_paragraph("网络层负责分组转发与路由选择。")
    doc.add_heading("4.1 IPv4 编址", level=2)
    doc.add_paragraph("CIDR 记法形如 192.168.1.0/24。")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "前缀长度"
    table.rows[0].cells[1].text = "可用主机数"
    table.rows[1].cells[0].text = "/24"
    table.rows[1].cells[1].text = "254"
    doc.save(str(path))


def _build_pdf(path: Path, pages: list[str], title: str = "") -> None:
    """手工拼一份最小合法 PDF（ASCII 文本、正确 xref），避免引入测试期 PDF 生成依赖。"""
    n = len(pages)
    page_ids = [3 + i for i in range(n)]
    content_ids = [3 + n + i for i in range(n)]
    font_id = 3 + 2 * n
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)

    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode(),
    ]
    for pid, cid in zip(page_ids, content_ids):
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {cid} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
    for text in pages:
        stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objs.append(b"<< /Title (" + title.encode() + b") >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R /Info {font_id + 1} 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    path.write_bytes(bytes(out))


def _make_sample_pptx(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "第4章 网络层"
    slide.placeholders[1].text = "分组转发与路由选择"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "4.1 IPv4 编址"
    body = slide.placeholders[1].text_frame
    body.text = "CIDR 记法形如 192.168.1.0/24"
    body.add_paragraph().text = "前缀长度决定可用主机数"
    slide.notes_slide.notes_text_frame.text = "课堂上强调前缀长度"

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "4.2 路由算法"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "算法"
    table.cell(0, 1).text = "类型"
    table.cell(1, 0).text = "距离向量"
    table.cell(1, 1).text = "分布式"

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "本页无标题占位符"

    prs.save(str(path))


async def test_text_parser_md(tmp_path):
    f = tmp_path / "第5章-传输层.md"
    f.write_text(MD_SAMPLE, encoding="utf-8")
    parsed = await get_parser(str(f)).parse(f)
    assert parsed.title == "第5章 传输层"
    assert "# 第5章 传输层" in parsed.markdown
    assert "UDP 是无连接的传输层协议" in parsed.markdown


async def test_text_parser_gbk_fallback(tmp_path):
    f = tmp_path / "legacy.txt"
    f.write_bytes("子网掩码与 CIDR 记法。".encode("gbk"))
    parsed = await get_parser(str(f)).parse(f)
    assert "子网掩码与 CIDR 记法" in parsed.markdown


async def test_docx_parser(tmp_path):
    f = tmp_path / "第4章-网络层.docx"
    _make_sample_docx(f)
    parsed = await get_parser(str(f)).parse(f)
    assert parsed.title == "第4章 网络层"
    assert "# 第4章 网络层" in parsed.markdown
    assert "## 4.1 IPv4 编址" in parsed.markdown
    assert "网络层负责分组转发与路由选择。" in parsed.markdown
    # 表格按文档顺序转 markdown，且在第二段之后
    assert "| 前缀长度 | 可用主机数 |" in parsed.markdown
    assert "| /24 | 254 |" in parsed.markdown
    assert parsed.markdown.index("4.1 IPv4 编址") < parsed.markdown.index("前缀长度")
    assert parsed.metadata["tables"] == 1


async def test_pdf_parser(tmp_path):
    f = tmp_path / "第4章-网络层.pdf"
    _build_pdf(
        f,
        ["Network Layer Routing", "CIDR notation 192.168.1.0/24"],
        title="Chapter 4 Network Layer",
    )
    parsed = await get_parser(str(f)).parse(f)
    assert parsed.title == "Chapter 4 Network Layer"
    assert "Network Layer Routing" in parsed.markdown
    assert "CIDR notation 192.168.1.0/24" in parsed.markdown
    assert parsed.markdown.index("Network Layer Routing") < parsed.markdown.index("CIDR notation")
    assert parsed.metadata["pages"] == 2


async def test_pdf_parser_title_falls_back_to_stem(tmp_path):
    f = tmp_path / "第4章-网络层.pdf"
    _build_pdf(f, ["Network Layer Routing"])
    parsed = await get_parser(str(f)).parse(f)
    assert parsed.title == "第4章-网络层"


async def test_pdf_scanned_like_raises(tmp_path):
    f = tmp_path / "scan.pdf"
    _build_pdf(f, [""])
    with pytest.raises(ValueError, match="未提取到文本"):
        await get_parser(str(f)).parse(f)


async def test_pptx_parser(tmp_path):
    f = tmp_path / "第4章-网络层.pptx"
    _make_sample_pptx(f)
    parsed = await get_parser(str(f)).parse(f)
    assert parsed.title == "第4章 网络层"
    assert "# 第4章 网络层" in parsed.markdown
    assert "## 4.1 IPv4 编址" in parsed.markdown
    assert "CIDR 记法形如 192.168.1.0/24" in parsed.markdown
    assert "前缀长度决定可用主机数" in parsed.markdown
    assert "备注：课堂上强调前缀长度" in parsed.markdown
    assert "| 算法 | 类型 |" in parsed.markdown
    assert "| 距离向量 | 分布式 |" in parsed.markdown
    assert "## 第 4 页" in parsed.markdown
    assert "本页无标题占位符" in parsed.markdown
    assert parsed.markdown.index("第4章 网络层") < parsed.markdown.index("4.1 IPv4 编址")
    assert parsed.metadata["slides"] == 4
    assert parsed.metadata["tables"] == 1


def test_unsupported_format(tmp_path):
    with pytest.raises(UnsupportedFormatError):
        get_parser("photo.png")


async def test_disabled_parser_raises_not_enabled(tmp_path, monkeypatch):
    class ReservedParser(BaseParser):
        extensions = [".reserved"]
        enabled = False

        async def parse(self, file_path: Path) -> ParsedDocument:
            raise NotEnabledError("预留格式")

    monkeypatch.setitem(PARSER_REGISTRY, ".reserved", ReservedParser)
    f = tmp_path / "draft.reserved"
    f.write_bytes(b"fake")
    with pytest.raises(NotEnabledError):
        await get_parser(str(f)).parse(f)
