import logging
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.parsers._markdown import table_to_markdown
from app.parsers.base import BaseParser, ParsedDocument, register

logger = logging.getLogger(__name__)


def _iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _table_markdown(table) -> str:
    return table_to_markdown(
        [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
    )


@register
class PptxParser(BaseParser):
    extensions = [".pptx"]

    async def parse(self, file_path: Path) -> ParsedDocument:
        try:
            prs = Presentation(str(file_path))
        except Exception as e:
            raise ValueError(f"pptx 文件无法打开（可能损坏或为加密文档）：{e}") from e

        title = file_path.stem
        title_seen = False
        has_content = False
        table_count = 0
        sections: list[str] = []

        for index, slide in enumerate(prs.slides, start=1):
            title_shape = slide.shapes.title
            slide_title = title_shape.text_frame.text.strip() if title_shape is not None else ""
            if slide_title:
                has_content = True
                if title_seen:
                    heading = f"## {slide_title}"
                else:
                    title_seen = True
                    title = slide_title
                    heading = f"# {slide_title}"
            else:
                heading = f"## 第 {index} 页"

            parts: list[str] = []
            for shape in _iter_shapes(slide.shapes):
                # shapes.title 每次迭代都会新建包装对象，只能按底层元素判等
                if title_shape is not None and shape._element is title_shape._element:
                    continue
                if shape.has_table:
                    table_count += 1
                    md = _table_markdown(shape.table)
                    if md:
                        parts.append(md)
                elif shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            parts.append(text)
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                notes = notes_frame.text.strip() if notes_frame is not None else ""
                if notes:
                    parts.append(f"备注：{notes}")
            if parts:
                has_content = True
                sections.append(heading + "\n\n" + "\n\n".join(parts))
            else:
                sections.append(heading)

        if not has_content:
            raise ValueError(
                "pptx 未提取到文本内容（可能整份为图片/图形）；请先转换为 md/txt/docx"
            )

        return ParsedDocument(
            title=title,
            markdown="\n\n".join(sections).strip(),
            metadata={"format": "pptx", "slides": len(prs.slides), "tables": table_count},
        )
