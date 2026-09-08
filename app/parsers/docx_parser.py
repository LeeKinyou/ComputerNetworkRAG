import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.parsers.base import BaseParser, ParsedDocument, register

_HEADING = re.compile(r"^(?:Heading|标题)\s*(\d+)\s*$", re.IGNORECASE)


def _iter_blocks(doc: DocumentObject):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _heading_level(paragraph: Paragraph) -> int | None:
    try:
        style_name = paragraph.style.name or ""
    except Exception:
        return None
    m = _HEADING.match(style_name.strip())
    return int(m.group(1)) if m else None


def _cell_text(cell) -> str:
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


def _table_markdown(table: Table) -> str:
    rows = [[_cell_text(c) for c in row.cells] for row in table.rows]
    rows = [r for r in rows if any(cell for cell in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |"]
    lines.append("| " + " | ".join(["---"] * width) + " |")
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


@register
class DocxParser(BaseParser):
    extensions = [".docx"]

    async def parse(self, file_path: Path) -> ParsedDocument:
        try:
            doc = Document(str(file_path))
        except Exception as e:
            raise ValueError(f"docx 文件无法打开（可能损坏或为加密文档）：{e}") from e

        parts: list[str] = []
        title = file_path.stem
        heading_count = 0
        table_count = 0

        for block in _iter_blocks(doc):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                level = _heading_level(block)
                if level:
                    heading_count += 1
                    if heading_count == 1:
                        title = text
                    parts.append("#" * min(level, 6) + " " + text)
                else:
                    parts.append(text)
            else:
                table_count += 1
                md = _table_markdown(block)
                if md:
                    parts.append(md)

        markdown = "\n\n".join(parts).strip()
        return ParsedDocument(
            title=title,
            markdown=markdown,
            metadata={"headings": heading_count, "tables": table_count},
        )
